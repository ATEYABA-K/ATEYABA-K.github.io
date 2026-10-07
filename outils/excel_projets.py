"""
Même traitement que les deux études, pour les 8 autres projets du portfolio :
pour chaque projet, projets/<projet>/excel/ contient un corrigé, un exercice, les CSV et les graphiques exportés.
Les données viennent des dépôts GitHub de chaque projet (clonés dans le dossier passé en argument).
Usage : python outils/excel_projets.py <dossier_des_depots>
"""
import csv
import sqlite3
import subprocess
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.series import SeriesLabel
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).parent))
from excel import (BLEU, F, GRIS, ORANGE, RECALC, a_faire, ecrire_csv, entetes, etiquettes,  # noqa: E402
                   exporter_graphiques, lisez_moi, onglet_graphique, point, remplir, styler, titre)

RACINE = Path(__file__).resolve().parent.parent
DEPOTS = Path(sys.argv[1]) if len(sys.argv) > 1 else RACINE.parent / "depots"
SANS = PatternFill(fill_type=None)


def entete_lisez_moi(wb, nom, exercice, onglets, source, outil=None):
    lisez_moi(wb, [f"{nom} · " + ("EXERCICE" if exercice else "CORRIGÉ"), "",
                   "Exercice : les cellules jaunes sont à remplir, les graphiques à construire. Suivez TUTO.md (ou la page Tuto du site)."
                   if exercice else "Corrigé : formules en place, et chaque graphique a son onglet G. Ce sont ces graphiques qui sont sur le site.",
                   *([outil] if outil else []), "", "Les onglets", *[f"• {o}" for o in onglets], "", "Source : " + source])


def barres(ws, col, r0, fin, titre_g, x=None, y=None, num=None, couleurs=None, horizontal=False, labels=None, cat_col=1):
    ch = BarChart()
    ch.gapWidth = 50
    if horizontal:
        ch.type = "bar"
        ch.x_axis.scaling.orientation = "maxMin"
    ch.add_data(Reference(ws, min_col=col, min_row=r0 - 1, max_row=fin), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=cat_col, min_row=r0, max_row=fin))
    s = ch.series[0]
    remplir(s, BLEU)
    s.invertIfNegative = False
    if couleurs:
        s.dPt += [point(k, c) for k, c in enumerate(couleurs)]
    if labels:
        etiquettes(s, labels)
    return styler(ch, titre_g, x, y, legende=False, num=num)


# =================================================================== Churn
def churn(exercice):
    rows = list(csv.DictReader(open(DEPOTS / "analyse-churn-telco/data/telco_churn.csv", encoding="utf-8")))
    wb = Workbook()
    entete_lisez_moi(wb, "Où se concentre le risque de départ client ?", exercice,
                     ["Clients : les 7 043 clients (ancienneté, internet, contrat, départ).", "Analyse : taux de départ par segment, avec NB.SI.ENS.",
                      "G1 à G3 : les graphiques."], "IBM Telco Customer Churn (jeu public), via le dépôt analyse-churn-telco.",
                     "Astuce : la même analyse se fait en 3 clics avec un tableau croisé dynamique (voir le tuto).")
    ws = wb.create_sheet("Clients")
    titre(ws, "Les clients", "Une ligne par client. Churn = Yes : le client est parti.",
          "À faire : la colonne F (tranche d'ancienneté), étape 1." if exercice else None)
    entetes(ws, 4, ["Client", "Ancienneté (mois)", "Internet", "Contrat", "Churn", "Tranche d'ancienneté"], {6: 20})
    n = len(rows)
    for i, r in enumerate(rows, 5):
        ws.cell(row=i, column=1, value=r["customerID"])
        ws.cell(row=i, column=2, value=int(r["tenure"]))
        ws.cell(row=i, column=3, value=r["InternetService"])
        ws.cell(row=i, column=4, value=r["Contract"])
        ws.cell(row=i, column=5, value=r["Churn"])
        if not exercice:
            ws.cell(row=i, column=6, value=f'=IF(B{i}<=12,"0-12 mois",IF(B{i}<=24,"13-24 mois",IF(B{i}<=48,"25-48 mois","49 mois et +")))')
    if exercice:
        a_faire(ws, "F5:F12")
    fin = 4 + n
    an = wb.create_sheet("Analyse")
    titre(an, "Taux de départ par segment", "Taux = clients partis du segment / clients du segment.",
          "À faire : les colonnes B, C, D des trois tableaux (étape 2), puis les graphiques (étape 3)." if exercice else None)
    blocs = [("Contrat", "D", ["Month-to-month", "One year", "Two year"]),
             ("Ancienneté", "F", ["0-12 mois", "13-24 mois", "25-48 mois", "49 mois et +"]),
             ("Internet", "C", ["Fiber optic", "DSL", "No"])]
    ligne, refs = 4, []
    for nom, col, cats in blocs:
        for j, h in enumerate([nom, "Clients", "Partis", "Taux de départ"], 1):
            c = an.cell(row=ligne, column=j, value=h)
            c.font, c.fill = F(bold=True, color="FFFFFF"), PatternFill("solid", fgColor=BLEU)
        r0 = ligne + 1
        for k, cat in enumerate(cats):
            i = r0 + k
            an.cell(row=i, column=1, value=cat)
            if not exercice:
                an[f"B{i}"] = f'=COUNTIF(Clients!${col}$5:${col}${fin},A{i})'
                an[f"C{i}"] = f'=COUNTIFS(Clients!${col}$5:${col}${fin},A{i},Clients!$E$5:$E${fin},"Yes")'
                an[f"D{i}"] = f"=C{i}/B{i}"
            an[f"D{i}"].number_format = "0.0%"
        if exercice:
            a_faire(an, f"B{r0}:D{r0 + len(cats) - 1}")
        refs.append((r0, r0 + len(cats) - 1))
        ligne = r0 + len(cats) + 2
    an[f"A{ligne}"], an[f"A{ligne}"].font = "Taux de départ global :", F(bold=True)
    if not exercice:
        an[f"D{ligne}"] = f'=COUNTIF(Clients!$E$5:$E${fin},"Yes")/COUNTA(Clients!$E$5:$E${fin})'
    else:
        a_faire(an, f"D{ligne}:D{ligne}")
    an[f"D{ligne}"].number_format = "0.0%"
    for c, w in zip("ABCD", (22, 12, 12, 16)):
        an.column_dimensions[c].width = w
    if not exercice:
        for (r0, r1), (g, t, cols) in zip(refs, [("G1 Contrat", "Taux de départ selon le contrat", [ORANGE, GRIS, BLEU]),
                                                  ("G2 Anciennete", "Taux de départ selon l'ancienneté", [ORANGE, GRIS, GRIS, BLEU]),
                                                  ("G3 Internet", "Taux de départ selon le type d'internet", [ORANGE, GRIS, BLEU])]):
            onglet_graphique(wb, g, barres(an, 4, r0, r1, t, None, "Taux de départ", "0%", cols, labels="0.0%"))
    ecrire_csv(OUT("churn-telco") / "csv", "clients.csv", ["client", "anciennete_mois", "internet", "contrat", "churn"],
               [[r["customerID"], r["tenure"], r["InternetService"], r["Contract"], r["Churn"]] for r in rows])
    return wb, {"G1 Contrat": "G1-contrat.png", "G2 Anciennete": "G2-anciennete.png", "G3 Internet": "G3-internet.png"}


# =================================================================== SQL
REQUETES = {
    "Genres": ("Chiffre d'affaires par genre (top 10)", """SELECT g.Name AS genre, ROUND(SUM(il.UnitPrice * il.Quantity), 2) AS ca_total
FROM InvoiceLine il
JOIN Track t ON t.TrackId = il.TrackId
JOIN Genre g ON g.GenreId = t.GenreId
GROUP BY g.Name ORDER BY ca_total DESC LIMIT 10;"""),
    "Cumul": ("Chiffre d'affaires cumulé mois par mois (fonction fenêtre)", """WITH ca_mensuel AS (
  SELECT strftime('%Y-%m', InvoiceDate) AS mois, ROUND(SUM(Total), 2) AS ca_du_mois
  FROM Invoice GROUP BY mois)
SELECT mois, ca_du_mois, ROUND(SUM(ca_du_mois) OVER (ORDER BY mois), 2) AS ca_cumule
FROM ca_mensuel ORDER BY mois;"""),
    "Clients": ("Top 10 des clients", """SELECT c.FirstName || ' ' || c.LastName AS client, c.Country AS pays, ROUND(SUM(i.Total), 2) AS ca_total
FROM Customer c JOIN Invoice i ON i.CustomerId = c.CustomerId
GROUP BY c.CustomerId ORDER BY ca_total DESC LIMIT 10;"""),
}


def sql(exercice):
    db = sqlite3.connect(DEPOTS / "analyse-sql-ventes-musique/chinook.db")
    wb = Workbook()
    entete_lisez_moi(wb, "Analyse SQL des ventes d'un magasin de musique", exercice,
                     ["Genres, Cumul, Clients : le résultat de chaque requête SQL, collé tel quel.", "Requetes : le texte des requêtes.", "G1 à G3 : les graphiques."],
                     "base Chinook (jeu public), via le dépôt analyse-sql-ventes-musique.",
                     "Outil : les requêtes se lancent dans DB Browser for SQLite (gratuit), puis on colle le résultat dans Excel. Voir le tuto.")
    total = db.execute("SELECT ROUND(SUM(Total), 2) FROM Invoice").fetchone()[0]
    out = {}
    for onglet, (t, q) in REQUETES.items():
        cur = db.execute(q)
        cols = [d[0] for d in cur.description]
        data = cur.fetchall()
        ecrire_csv(OUT("sql-musique") / "csv", f"{onglet.lower()}.csv", cols, data)
        ws = wb.create_sheet(onglet)
        titre(ws, t, "Résultat de la requête, exporté de DB Browser for SQLite (Fichier → Exporter → Table vers CSV).",
              "À faire : coller le résultat de la requête (étape 1), puis le graphique (étape 2)." if exercice else None)
        entetes(ws, 4, cols, {1: 22})
        for i, r in enumerate(data, 5):
            for j, v in enumerate(r, 1):
                c = ws.cell(row=i, column=j, value=None if exercice else v)
                if isinstance(v, float):
                    c.number_format = "#,##0.00"
        if exercice:
            a_faire(ws, f"A5:{get_column_letter(len(cols))}{4 + len(data)}")
        out[onglet] = (ws, 5, 4 + len(data))
    ws, r0, fin = out["Genres"]
    ws["D4"], ws["D4"].font = "Part du CA total", F(bold=True)
    ws["F2"], ws["F2"].font = "CA total du magasin :", F(bold=True)
    ws["G2"], ws["G2"].number_format = (None if exercice else total), "#,##0.00"
    for i in range(r0, fin + 1):
        if not exercice:
            ws[f"D{i}"] = f"=B{i}/$G$2"
        ws[f"D{i}"].number_format = "0.0%"
    if exercice:
        a_faire(ws, f"D{r0}:D{fin}")
    rq = wb.create_sheet("Requetes")
    titre(rq, "Les requêtes SQL", "À copier dans l'onglet « Exécuter le SQL » de DB Browser for SQLite.")
    k = 4
    for onglet, (t, q) in REQUETES.items():
        rq.cell(row=k, column=1, value=f"-- {t}").font = F(bold=True)
        for ligne in q.splitlines():
            k += 1
            rq.cell(row=k, column=1, value=ligne).font = Font(name="Courier New")
        k += 2
    rq.column_dimensions["A"].width = 110
    if not exercice:
        ws, r0, fin = out["Genres"]
        onglet_graphique(wb, "G1 Genres", barres(ws, 2, r0, fin, "Chiffre d'affaires par genre musical (top 10)", None, "CA (€)", "0",
                                                 [ORANGE] + [GRIS] * 9, horizontal=True, labels="#,##0"))
        ws, r0, fin = out["Cumul"]
        lc = LineChart()
        lc.add_data(Reference(ws, min_col=3, min_row=4, max_row=fin), titles_from_data=True)
        lc.set_categories(Reference(ws, min_col=1, min_row=r0, max_row=fin))
        lc.series[0].graphicalProperties.line.solidFill, lc.series[0].graphicalProperties.line.width = BLEU, 31750
        lc.series[0].smooth = False
        lc.x_axis.tickLblSkip = 6
        onglet_graphique(wb, "G2 Cumul", styler(lc, "Chiffre d'affaires cumulé, mois par mois", "Mois", "CA cumulé (€)", legende=False, num="#,##0"))
        ws, r0, fin = out["Clients"]
        ch = barres(ws, 3, r0, fin, "Top 10 des clients par chiffre d'affaires", None, "CA (€)", "0", [BLEU] * 10, horizontal=True, labels="#,##0.00")
        ch.y_axis.scaling.min = 0  # un axe qui ne part pas de 0 exagère les écarts
        onglet_graphique(wb, "G3 Clients", ch)
    return wb, {"G1 Genres": "G1-genres.png", "G2 Cumul": "G2-cumul.png", "G3 Clients": "G3-clients.png"}


# =================================================================== Scoring
def scoring(exercice):
    rows = list(csv.DictReader(open(DEPOTS / "scoring-prospects-b2b/data/resultats_deciles.csv")))
    wb = Workbook()
    entete_lisez_moi(wb, "Scoring de propension B2B : qui appeler en premier ?", exercice,
                     ["Deciles : le taux de conversion de chaque groupe de 10 %, du mieux noté au moins bien noté.", "G1 : le graphique du lift."],
                     "UCI Bank Marketing (jeu public), via le dépôt scoring-prospects-b2b. Le modèle (scikit-learn) est dans le dépôt.")
    ws = wb.create_sheet("Deciles")
    titre(ws, "Taux de conversion par décile de score", "Décile 1 = les 10 % de prospects les mieux notés par le modèle. Lift = taux du décile / taux moyen.",
          "À faire : le taux moyen (E2), la colonne D (lift), étape 1, puis le graphique, étape 2." if exercice else None)
    entetes(ws, 4, ["Décile", "Taux de conversion", "Prospects", "Lift"])
    for i, r in enumerate(rows, 5):
        ws.cell(row=i, column=1, value=f"D{int(r['decile']) + 1}")
        ws.cell(row=i, column=2, value=round(float(r["taux_conversion"]), 4)).number_format = "0.0%"
        ws.cell(row=i, column=3, value=int(r["n"])).number_format = "#,##0"
        ws[f"D{i}"].number_format = "0.0"
        if not exercice:
            ws[f"D{i}"] = f"=B{i}/$E$2"
    fin = 4 + len(rows)
    ws["D2"], ws["D2"].font = "Taux moyen :", F(bold=True)
    ws["E2"].number_format = "0.0%"
    if exercice:
        a_faire(ws, "E2:E2"); a_faire(ws, f"D5:D{fin}")
    else:
        ws["E2"] = f"=SUMPRODUCT(B5:B{fin},C5:C{fin})/SUM(C5:C{fin})"
        cols = [BLEU if float(r["lift"]) > 1 else GRIS for r in rows]
        onglet_graphique(wb, "G1 Lift", barres(ws, 4, 5, fin, "Lift par décile : au-dessus de 1, mieux que la moyenne", "Décile de score (D1 = 10 % les mieux notés)",
                                               "Lift", "0", cols, labels="0.0"))
    ecrire_csv(OUT("scoring-b2b") / "csv", "deciles.csv", ["decile", "taux_conversion", "prospects"],
               [[f"D{int(r['decile']) + 1}", r["taux_conversion"], r["n"]] for r in rows])
    return wb, {"G1 Lift": "G1-lift.png"}


# =================================================================== Insee IA (2 projets)
def insee_ia(exercice, projet):
    rows = [r for r in csv.DictReader(open(DEPOTS / "adoption-ia-pme-france/data/insee_ia_pour_powerbi.csv", encoding="utf-8-sig"))
            if r["type"] == "Taille de l'entreprise" and r["zone"] == "France"]
    tailles = ["De 10 à 49 salariés", "De 50 à 249 salariés", "250 salariés ou plus"]
    annees = ["2023", "2024", "2025"]
    val = {(r["categorie"].strip(), r["annee"]): float(r["part_entreprises_ia_pct"]) / 100 for r in rows}
    ia = projet == "ia-pme"
    wb = Workbook()
    entete_lisez_moi(wb, "Adoption de l'IA par les entreprises françaises" if ia else "Automatisation d'un reporting Insee", exercice,
                     ["Donnees : part des entreprises utilisant au moins une technologie d'IA, par taille et par année.", "G1 : le graphique."],
                     "Insee Première n° 2120, enquête TIC 2023-2025.",
                     "Le même graphique existe en tableau de bord Looker Studio (lien dans le tuto)." if ia
                     else "Le fichier Insee est téléchargé et nettoyé chaque mois par un robot (GitHub Actions) : le CSV est ensuite prêt pour Excel.")
    ws = wb.create_sheet("Donnees")
    titre(ws, "Part des entreprises qui utilisent l'IA", "Source : Insee, enquête TIC. En % des entreprises de chaque taille.",
          "À faire : la ligne Écart (étape 1), puis le graphique (étape 2)." if exercice else None)
    if ia:
        entetes(ws, 4, ["Taille", *annees], {1: 24})
        for i, t in enumerate(tailles, 5):
            ws.cell(row=i, column=1, value=t)
            for j, a in enumerate(annees, 2):
                ws.cell(row=i, column=j, value=val[(t, a)]).number_format = "0%"
        ws["A9"], ws["A9"].font = "Écart 250+ vs 10-49 (points)", F(bold=True)
        for j in range(2, 5):
            col = get_column_letter(j)
            ws[f"{col}9"].number_format = "0"
            if not exercice:
                ws[f"{col}9"] = f"=({col}7-{col}5)*100"
        if exercice:
            a_faire(ws, "B9:D9")
        else:
            ch = BarChart(); ch.gapWidth = 60
            ch.add_data(Reference(ws, min_col=2, max_col=4, min_row=4, max_row=7), titles_from_data=True)
            ch.set_categories(Reference(ws, min_col=1, min_row=5, max_row=7))
            for s, c in zip(ch.series, (GRIS, "4A5F8A", BLEU)):
                remplir(s, c)
                etiquettes(s, "0%")
            onglet_graphique(wb, "G1 Adoption", styler(ch, "Part des entreprises qui utilisent l'IA, par taille", None, None, num="0%"))
        graph = {"G1 Adoption": "G1-adoption.png"}
    else:
        entetes(ws, 4, ["Année", *tailles], {2: 20, 3: 20, 4: 20})
        for i, a in enumerate(annees, 5):
            ws.cell(row=i, column=1, value=a)
            for j, t in enumerate(tailles, 2):
                ws.cell(row=i, column=j, value=val[(t, a)]).number_format = "0%"
        ws["A9"], ws["A9"].font = "Progression 2023 → 2025 (points)", F(bold=True)
        for j in range(2, 5):
            col = get_column_letter(j)
            ws[f"{col}9"].number_format = "0"
            if not exercice:
                ws[f"{col}9"] = f"=({col}7-{col}5)*100"
        if exercice:
            a_faire(ws, "B9:D9")
        else:
            lc = LineChart()
            lc.add_data(Reference(ws, min_col=2, max_col=4, min_row=4, max_row=7), titles_from_data=True)
            lc.set_categories(Reference(ws, min_col=1, min_row=5, max_row=7))
            for s, c in zip(lc.series, (GRIS, "4A5F8A", ORANGE)):
                s.graphicalProperties.line.solidFill, s.graphicalProperties.line.width = c, 34925
                s.smooth = False
                etiquettes(s, "0%")
            onglet_graphique(wb, "G1 Evolution", styler(lc, "Adoption de l'IA de 2023 à 2025, par taille d'entreprise", "Année", None, num="0%"))
        graph = {"G1 Evolution": "G1-evolution.png"}
    ecrire_csv(OUT(projet) / "csv", "adoption_ia.csv", ["taille", "annee", "part"],
               [[t, a, val[(t, a)]] for t in tailles for a in annees])
    return wb, graph


# =================================================================== IPC
def lire_insee(fichier):
    out = {}
    for r in csv.reader(open(fichier, encoding="utf-8-sig"), delimiter=";"):
        if r and len(r[0]) == 7 and r[0][4] == "-":
            try:
                out[r[0]] = float(r[1].replace(",", "."))
            except ValueError:
                pass
    return out


def ipc(exercice):
    D = DEPOTS / "analyse-ipc-france-covid/data"
    g, a, e, s = (lire_insee(D / f) for f in ("ipc_global.csv", "ipc_alimentation.csv", "ipc_energie.csv", "smic_brut_mensuel.csv"))
    mois = sorted(m for m in g if "2019-01" <= m and m in a and m in e and m in s)
    wb = Workbook()
    entete_lisez_moi(wb, "Évolution des prix à la consommation en France", exercice,
                     ["Donnees : les indices Insee (base 2025) et le Smic brut, chaque mois depuis janvier 2019.",
                      "Colonnes F à I : les mêmes indices rebasés à 100 en janvier 2020, et le pouvoir d'achat du Smic.", "G1, G2 : les graphiques."],
                     "Insee, indices des prix à la consommation (idBank 011814131, 011813717, 011813867) et Smic brut (000879877).")
    ws = wb.create_sheet("Donnees")
    titre(ws, "Prix et Smic, rebasés à 100 en janvier 2020", "Rebaser : diviser chaque valeur par celle de janvier 2020, multiplier par 100.",
          "À faire : colonnes F à I (étapes 1 et 2), puis les graphiques (étape 3)." if exercice else None)
    entetes(ws, 4, ["Mois", "IPC global", "IPC alimentation", "IPC énergie", "Smic brut (€)",
                    "Global (base 100)", "Alimentation (base 100)", "Énergie (base 100)", "Pouvoir d'achat du Smic"], {1: 10})
    ref = 5 + mois.index("2020-01")
    for i, m in enumerate(mois, 5):
        ws.cell(row=i, column=1, value=m)
        for j, serie in enumerate((g, a, e, s), 2):
            ws.cell(row=i, column=j, value=serie[m]).number_format = "0.00"
        for col in "FGHI":
            ws[f"{col}{i}"].number_format = "0.0"
        if not exercice:
            ws[f"F{i}"], ws[f"G{i}"], ws[f"H{i}"] = f"=B{i}/B${ref}*100", f"=C{i}/C${ref}*100", f"=D{i}/D${ref}*100"
            ws[f"I{i}"] = f"=(E{i}/E${ref})/(B{i}/B${ref})*100"
    fin = 4 + len(mois)
    ws.cell(row=ref, column=1).fill = PatternFill("solid", fgColor="DDEBF7")
    if exercice:
        a_faire(ws, f"F5:I{fin}")
    else:
        lc = LineChart()
        lc.add_data(Reference(ws, min_col=6, max_col=8, min_row=4, max_row=fin), titles_from_data=True)
        lc.set_categories(Reference(ws, min_col=1, min_row=5, max_row=fin))
        for s_, c in zip(lc.series, (BLEU, GRIS, ORANGE)):
            s_.graphicalProperties.line.solidFill, s_.graphicalProperties.line.width = c, 28575
            s_.smooth = False
        lc.x_axis.tickLblSkip = 12
        lc.y_axis.scaling.min = 80
        onglet_graphique(wb, "G1 Prix", styler(lc, "Prix à la consommation depuis 2019 (base 100 = janvier 2020)", "Mois", "Indice", num="0"))
        l2 = LineChart()
        l2.add_data(Reference(ws, min_col=9, min_row=4, max_row=fin), titles_from_data=True)
        l2.set_categories(Reference(ws, min_col=1, min_row=5, max_row=fin))
        l2.series[0].graphicalProperties.line.solidFill, l2.series[0].graphicalProperties.line.width = ORANGE, 31750
        l2.series[0].smooth = False
        l2.x_axis.tickLblSkip = 12
        l2.y_axis.scaling.min, l2.y_axis.scaling.max = 96, 106
        onglet_graphique(wb, "G2 Smic", styler(l2, "Pouvoir d'achat du Smic (base 100 = janvier 2020)", "Mois", "Indice", legende=False, num="0"))
    ecrire_csv(OUT("ipc-france") / "csv", "ipc_smic.csv", ["mois", "ipc_global", "ipc_alimentation", "ipc_energie", "smic_brut"],
               [[m, g[m], a[m], e[m], s[m]] for m in mois])
    return wb, {"G1 Prix": "G1-prix.png", "G2 Smic": "G2-smic.png"}


# =================================================================== Cadrage : Gantt
JALONS = [("Cadrage validé", 0, 0), ("Données consolidées", 1, 3), ("Premier modèle", 4, 6),
          ("Validation métier", 7, 9), ("Intégration CRM", 10, 12), ("Bilan à 3 mois", 24, 24)]


def cadrage(exercice):
    wb = Workbook()
    entete_lisez_moi(wb, "Note de cadrage : le planning en diagramme de Gantt", exercice,
                     ["Jalons : les 6 jalons de la note de cadrage, semaine de début et de fin.", "G1 : le diagramme de Gantt."],
                     "note de cadrage (dépôt note-cadrage-scoring-ia-pme), cas fictif.")
    ws = wb.create_sheet("Jalons")
    titre(ws, "Planning des jalons", "Un Gantt dans Excel = un histogramme empilé dont la première série est invisible.",
          "À faire : la colonne D (durée), étape 1, puis le Gantt, étape 2." if exercice else None)
    entetes(ws, 4, ["Jalon", "Semaine de début", "Semaine de fin", "Durée (semaines)"], {1: 24})
    for i, (j, d, f) in enumerate(JALONS, 5):
        ws.cell(row=i, column=1, value=j); ws.cell(row=i, column=2, value=d); ws.cell(row=i, column=3, value=f)
        if not exercice:
            ws[f"D{i}"] = f"=C{i}-B{i}+1"
    fin = 4 + len(JALONS)
    if exercice:
        a_faire(ws, f"D5:D{fin}")
    else:
        ch = BarChart()
        ch.type, ch.grouping, ch.overlap, ch.gapWidth = "bar", "stacked", 100, 40
        ch.add_data(Reference(ws, min_col=2, min_row=4, max_row=fin), titles_from_data=True)
        ch.add_data(Reference(ws, min_col=4, min_row=4, max_row=fin), titles_from_data=True)
        ch.set_categories(Reference(ws, min_col=1, min_row=5, max_row=fin))
        ch.series[0].graphicalProperties.noFill = True
        ch.series[0].graphicalProperties.line.noFill = True
        remplir(ch.series[1], BLEU)
        ch.series[1].dPt.append(point(len(JALONS) - 1, ORANGE))
        ch.x_axis.scaling.orientation = "maxMin"
        ch.y_axis.scaling.min, ch.y_axis.scaling.max = 0, 26
        ch.y_axis.majorUnit = 2
        onglet_graphique(wb, "G1 Gantt", styler(ch, "Planning du projet de scoring (semaines)", None, "Semaine", legende=False, num="0"))
    ecrire_csv(OUT("cadrage-gantt") / "csv", "jalons.csv", ["jalon", "semaine_debut", "semaine_fin"], JALONS)
    return wb, {"G1 Gantt": "G1-gantt.png"}


# =================================================================== Suivi de projet
def suivi(exercice):
    src = load_workbook(DEPOTS / "suivi-projet-dashboard-kpi/suivi_projet_scoring_churn.xlsx")["Suivi_Jalons"]
    jal = [r for r in src.iter_rows(min_row=2, max_row=7, values_only=True)]
    wb = Workbook()
    entete_lisez_moi(wb, "Suivre un projet après l'avoir cadré (semaine 14)", exercice,
                     ["Suivi : les 6 jalons, leur statut, le budget prévu et réel.", "Indicateurs : avancement, retards, écart budgétaire (formules).", "G1 : budget prévu vs réel."],
                     "dépôt suivi-projet-dashboard-kpi, cas fictif.")
    ws = wb.create_sheet("Suivi")
    titre(ws, "Suivi des jalons à la semaine 14", "Budget réel = dépensé à ce jour.",
          "À faire : la colonne E (écart) et les indicateurs (étape 1), puis le graphique (étape 2)." if exercice else None)
    entetes(ws, 4, ["Jalon", "Statut", "Budget prévu (€)", "Budget réel (€)", "Écart (€)"], {1: 24, 2: 12})
    for i, r in enumerate(jal, 5):
        ws.cell(row=i, column=1, value=r[0]); ws.cell(row=i, column=2, value=r[4])
        ws.cell(row=i, column=3, value=r[5]).number_format = "#,##0"
        ws.cell(row=i, column=4, value=r[6]).number_format = "#,##0"
        ws[f"E{i}"].number_format = "+#,##0;-#,##0;0"
        if not exercice:
            ws[f"E{i}"] = f"=IF(D{i}=0,0,D{i}-C{i})"
    fin = 4 + len(jal)
    ind = [("Jalons terminés", f'=COUNTIF(B5:B{fin},"Terminé")', "0"),
           ("Jalons en retard", f'=COUNTIF(B5:B{fin},"En retard")', "0"),
           ("Avancement", f'=COUNTIF(B5:B{fin},"Terminé")/COUNTA(B5:B{fin})', "0%"),
           ("Budget prévu des jalons engagés (€)", f'=SUMIF(D5:D{fin},">0",C5:C{fin})', "#,##0"),
           ("Budget réel (€)", f"=SUM(D5:D{fin})", "#,##0"),
           ("Écart budgétaire", "=E18/E17-1", "+0.0%;-0.0%")]
    ws["A13"], ws["A13"].font = "Indicateurs", F(bold=True, size=12)
    for k, (lab, f, fmt) in enumerate(ind):
        ws[f"A{14 + k}"] = lab
        ws[f"E{14 + k}"].number_format = fmt
        if not exercice:
            ws[f"E{14 + k}"] = f
    if exercice:
        a_faire(ws, f"E5:E{fin}"); a_faire(ws, "E14:E19")
    else:
        ch = BarChart(); ch.gapWidth = 60
        ch.add_data(Reference(ws, min_col=3, max_col=4, min_row=4, max_row=fin), titles_from_data=True)
        ch.set_categories(Reference(ws, min_col=1, min_row=5, max_row=fin))
        remplir(ch.series[0], GRIS); remplir(ch.series[1], BLEU)
        ch.series[1].dPt.append(point(1, ORANGE))
        onglet_graphique(wb, "G1 Budget", styler(ch, "Budget prévu et réel par jalon, semaine 14 (orange : dépassement)", None, "€", num="#,##0"))
    ecrire_csv(OUT("suivi-projet") / "csv", "suivi_jalons.csv", ["jalon", "statut", "budget_prevu", "budget_reel"],
               [[r[0], r[4], r[5], r[6]] for r in jal])
    return wb, {"G1 Budget": "G1-budget.png"}


def OUT(projet):
    return RACINE / "projets" / projet / "excel"


PROJETS = {
    "churn-telco": churn, "sql-musique": sql, "scoring-b2b": scoring,
    "ia-pme": lambda exo: insee_ia(exo, "ia-pme"), "reporting-insee": lambda exo: insee_ia(exo, "reporting-insee"),
    "ipc-france": ipc, "cadrage-gantt": cadrage, "suivi-projet": suivi,
}

if __name__ == "__main__":
    seuls = sys.argv[2:] or list(PROJETS)
    for nom in seuls:
        fab = PROJETS[nom]
        dossier = OUT(nom)
        dossier.mkdir(parents=True, exist_ok=True)
        for exo in (False, True):
            wb, graphes = fab(exo)
            f = dossier / f"{nom}-{'exercice' if exo else 'corrige'}.xlsx"
            wb.save(f)
            if RECALC.exists():
                r = subprocess.run(["python3", str(RECALC), str(f), "180"], capture_output=True, text=True).stdout
                print(f.name, "erreurs:", r.split('"total_errors": ')[1].split(",")[0] if "total_errors" in r else r[:200])
        exporter_graphiques(dossier / f"{nom}-corrige.xlsx", dossier / "graphiques", graphes)
        print(nom, "ok")
