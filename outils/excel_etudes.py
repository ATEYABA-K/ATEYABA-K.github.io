"""Classeurs Excel des études « Sports de combat en France » et « Messi », et export des graphiques en PNG.

Utilisation : python outils/excel_etudes.py
Les classeurs restent en local (voir .gitignore), seuls les graphiques exportés sont publiés.
"""
import json
import subprocess

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment

from excel import (BLEU, F, GRIS, ORANGE, RACINE, RECALC, entetes, etiquettes, exporter_graphiques, lisez_moi,
                   onglet_graphique, point, remplir, styler, titre)


def nombre(ws, plage, fmt):
    for ligne in ws[plage]:
        for c in ligne:
            c.number_format = fmt


# ============================================================== Sports de combat
def combat():
    d = json.loads((RACINE / "projets/sports-combat/donnees.json").read_text())
    wb = Workbook()
    lisez_moi(wb, ["Sports de combat en France : croissance, pics et stratégies",
                   "Données compilées à partir de sources publiques (Insee, fédérations, organisateurs). Chaque ligne a sa source.",
                   "", "Les onglets",
                   "Taille : licences 2024 par discipline.", "Croissance : boxe et MMA, saison par saison.",
                   "Judo : licences pour 10 000 habitants, 2019 à 2024.", "Salles : spectateurs des grands événements MMA à Paris."])

    # Taille du marché
    ws = wb.create_sheet("Taille")
    titre(ws, "Licences par discipline (2024 ou dernière saison connue)", "Sources : Insee, FFBoxe, FFSavate, IMMAF")
    entetes(ws, 4, ["Discipline", "Licences", "Source"], {1: 26, 2: 14, 3: 70})
    lignes = sorted(d["taille_2024"], key=lambda x: x["licences"])
    for i, x in enumerate(lignes, 5):
        ws.cell(row=i, column=1, value=x["sport"]); ws.cell(row=i, column=2, value=x["licences"]); ws.cell(row=i, column=3, value=x["source"])
    fin = 4 + len(lignes)
    nombre(ws, f"B5:B{fin}", "# ##0")
    ch = BarChart(); ch.type = "bar"; ch.gapWidth = 60
    ch.add_data(Reference(ws, min_col=2, min_row=4, max_row=fin), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=5, max_row=fin))
    styler(ch, "Combien de licenciés ? Le MMA reste tout petit", legende=False, num="# ##0")
    s = ch.series[0]; remplir(s, GRIS); etiquettes(s, "# ##0")
    s.dPt = [point(lignes.index(next(x for x in lignes if x["sport"] == "MMA")), ORANGE)]
    onglet_graphique(wb, "G1 Taille", ch)

    # Croissance boxe et MMA
    ws = wb.create_sheet("Croissance")
    titre(ws, "Licences de boxe et de MMA, saison par saison", "Sources : FFBoxe, Insee, IMMAF, Les Pharaons")
    entetes(ws, 4, ["Période", "Licences", "Discipline", "Note", "Source"], {1: 12, 2: 12, 3: 12, 4: 52, 5: 70})
    rows = [(x["periode"], x["valeur"], "Boxe", x["note"], x["source"]) for x in d["licences_boxe"]] + \
           [(x["periode"], x["valeur"], "MMA", x["note"], x["source"]) for x in d["licences_mma"]]
    for i, r in enumerate(rows, 5):
        for j, v in enumerate(r, 1):
            ws.cell(row=i, column=j, value=v)
    nombre(ws, f"B5:B{4 + len(rows)}", "# ##0")
    # Indice base 100 à la première saison connue, pour comparer deux tailles très différentes.
    entetes(ws, 13, ["Période", "Boxe (base 100)", "MMA (base 100)"], {1: 12, 2: 16, 3: 16})
    periodes = ["2020-21", "2021-22", "2022", "2023-24", "2024-25"]
    boxe = {x["periode"]: x["valeur"] for x in d["licences_boxe"]}
    mma = {x["periode"]: x["valeur"] for x in d["licences_mma"]}
    for i, p in enumerate(periodes, 14):
        ws.cell(row=i, column=1, value=p)
        if p in boxe:
            ws.cell(row=i, column=2, value=f"=ROUND({boxe[p]}/{boxe['2020-21']}*100,0)")
        if p in mma:
            ws.cell(row=i, column=3, value=f"=ROUND({mma[p]}/{mma['2021-22']}*100,0)")
    ch = BarChart(); ch.gapWidth = 50
    ch.add_data(Reference(ws, min_col=2, max_col=3, min_row=13, max_row=18), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=14, max_row=18))
    styler(ch, "Croissance des licences (100 = première saison connue)", y="Indice")
    remplir(ch.series[0], BLEU); remplir(ch.series[1], ORANGE)
    for s in ch.series:
        etiquettes(s)
    onglet_graphique(wb, "G2 Croissance", ch)

    # Judo
    ws = wb.create_sheet("Judo")
    titre(ws, "Judo : licences pour 10 000 habitants", "Source : Insee, série 001756153")
    entetes(ws, 4, ["Année", "Licences pour 10 000 hab."], {1: 10, 2: 24})
    ans = [k for k in d["judo_pour_10000"] if k.isdigit()]
    for i, a in enumerate(ans, 5):
        ws.cell(row=i, column=1, value=a); ws.cell(row=i, column=2, value=d["judo_pour_10000"][a])
    ch = LineChart()
    ch.add_data(Reference(ws, min_col=2, min_row=4, max_row=4 + len(ans)), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=5, max_row=4 + len(ans)))
    styler(ch, "Judo : le creux du Covid, puis la reprise jusqu'aux Jeux", legende=False)
    s = ch.series[0]; s.graphicalProperties.line.solidFill = BLEU; s.graphicalProperties.line.width = 32000
    s.smooth = False; etiquettes(s); s.dLbls.position = "t"
    ch.y_axis.scaling.min = 0
    onglet_graphique(wb, "G3 Judo", ch)

    # Salles
    ws = wb.create_sheet("Salles")
    titre(ws, "Spectateurs des grands événements MMA à Paris", "Sources : UFC, Sortiraparis")
    entetes(ws, 4, ["Événement", "Spectateurs", "Source"], {1: 30, 2: 14, 3: 70})
    for i, x in enumerate(d["salles"], 5):
        ws.cell(row=i, column=1, value=x["evenement"]); ws.cell(row=i, column=2, value=x["spectateurs"]); ws.cell(row=i, column=3, value=x["source"])
    nombre(ws, "B5:B7", "# ##0")
    return wb


# ============================================================== Messi
def messi():
    r = json.loads((RACINE / "projets/messi/resultats.json").read_text())
    wb = Workbook()
    lisez_moi(wb, ["Messi est-il le meilleur joueur de l'histoire ?",
                   "Sources : Transfermarkt (octobre 2025, jeu de données salimt/football-datasets) et StatsBomb (données ouvertes, Liga 2004-05 à 2020-21).",
                   "", "Les onglets",
                   "Carriere : buts, penalties, passes et matchs de 14 grands attaquants.",
                   "Saisons : buts réels et buts attendus (xG) de Messi en Liga, saison par saison.",
                   "Ballons : Ballon d'Or remportés."])

    c = r["carriere"]
    ws = wb.create_sheet("Carriere")
    titre(ws, "Carrière en club et en sélection", "Club : toutes compétitions. Sélection : équipe A.")
    cols = ["Joueur", "Matchs club", "Buts club", "Penalties", "Passes déc.", "Buts sélection", "Buts au total",
            "Buts hors penalty", "(Buts hors pen. + passes) par match", "Saisons à 40 buts"]
    entetes(ws, 4, cols, {1: 20, 9: 20})
    for i, x in enumerate(c, 5):
        vals = [x["nom"], x["matchs"], x["buts"], x["penalties"], x["passes"], x["buts_sel"]]
        for j, v in enumerate(vals, 1):
            ws.cell(row=i, column=j, value=v)
        ws.cell(row=i, column=7, value=f"=C{i}+F{i}")
        ws.cell(row=i, column=8, value=f"=C{i}-D{i}")
        ws.cell(row=i, column=9, value=f"=ROUND((H{i}+E{i})/B{i},2)")
        ws.cell(row=i, column=10, value=x["saisons_40"])
    fin = 4 + len(c)

    # Tri pour les graphiques : tableau des totaux (buts au total), classé.
    tot = sorted(c, key=lambda x: x["buts_total"])
    ws2 = wb.create_sheet("Totaux")
    titre(ws2, "Buts au total (club + sélection)", "Classement par nombre de buts")
    entetes(ws2, 4, ["Joueur", "Buts au total", "Contribution par match"], {1: 20, 2: 14, 3: 22})
    for i, x in enumerate(tot, 5):
        ws2.cell(row=i, column=1, value=x["nom"]); ws2.cell(row=i, column=2, value=int(x["buts_total"]))
    ch = BarChart(); ch.type = "bar"; ch.gapWidth = 50
    ch.add_data(Reference(ws2, min_col=2, min_row=4, max_row=4 + len(tot)), titles_from_data=True)
    ch.set_categories(Reference(ws2, min_col=1, min_row=5, max_row=4 + len(tot)))
    styler(ch, "Buts au total : Cristiano Ronaldo devant", legende=False, num="0")
    s = ch.series[0]; remplir(s, GRIS); etiquettes(s, "0")
    s.dPt = [point([x["nom"] for x in tot].index("Messi"), ORANGE), point([x["nom"] for x in tot].index("Cristiano Ronaldo"), BLEU)]
    onglet_graphique(wb, "G1 Totaux", ch)

    par = sorted(c, key=lambda x: x["contrib_par_match"])
    ws3 = wb.create_sheet("Par match")
    titre(ws3, "Buts hors penalty + passes décisives, par match de club", "Ce qu'un joueur apporte vraiment à chaque match")
    entetes(ws3, 4, ["Joueur", "Contribution par match"], {1: 20, 2: 22})
    for i, x in enumerate(par, 5):
        ws3.cell(row=i, column=1, value=x["nom"]); ws3.cell(row=i, column=2, value=round(x["contrib_par_match"], 2))
    ch = BarChart(); ch.type = "bar"; ch.gapWidth = 50
    ch.add_data(Reference(ws3, min_col=2, min_row=4, max_row=4 + len(par)), titles_from_data=True)
    ch.set_categories(Reference(ws3, min_col=1, min_row=5, max_row=4 + len(par)))
    styler(ch, "Sans penalty, avec les passes : Messi devant tout le monde", legende=False, num="0.00")
    s = ch.series[0]; remplir(s, GRIS); etiquettes(s, "0.00")
    s.dPt = [point([x["nom"] for x in par].index("Messi"), ORANGE), point([x["nom"] for x in par].index("Cristiano Ronaldo"), BLEU)]
    onglet_graphique(wb, "G2 Par match", ch)

    sa = r["tirs"]["saisons"]
    ws4 = wb.create_sheet("Saisons")
    titre(ws4, "Messi en Liga : buts réels et buts attendus (hors penalty)", "xG : probabilité de but de chaque tir, selon StatsBomb")
    entetes(ws4, 4, ["Saison", "Tirs", "Buts", "Buts attendus (xG)", "Écart"], {1: 12, 4: 18})
    for i, x in enumerate(sa, 5):
        ws4.cell(row=i, column=1, value=x["saison"][2:4] + "-" + x["saison"][7:9]); ws4.cell(row=i, column=2, value=x["tirs"])
        ws4.cell(row=i, column=3, value=int(x["buts"])); ws4.cell(row=i, column=4, value=round(x["xg"], 1))
        ws4.cell(row=i, column=5, value=f"=C{i}-D{i}")
    f4 = 4 + len(sa)
    nombre(ws4, f"D5:E{f4}", "0.0")
    ch = BarChart(); ch.gapWidth = 40
    ch.add_data(Reference(ws4, min_col=3, max_col=4, min_row=4, max_row=f4), titles_from_data=True)
    ch.set_categories(Reference(ws4, min_col=1, min_row=5, max_row=f4))
    styler(ch, "16 saisons sur 17, plus de buts que ses occasions n'en valaient", y="Buts")
    remplir(ch.series[0], ORANGE); remplir(ch.series[1], GRIS)
    ch.x_axis.txPr = __import__("excel").texte(800)
    onglet_graphique(wb, "G3 xG", ch)

    b = r["ballons"]
    ws5 = wb.create_sheet("Ballons")
    titre(ws5, "Ballon d'Or remportés", "Source : France Football, liste publiée par Sports Illustrated")
    entetes(ws5, 4, ["Joueur", "Ballons d'Or", "Années"], {1: 22, 2: 12, 3: 52})
    lb = sorted(b.items(), key=lambda kv: len(kv[1]))
    for i, (nom, ans) in enumerate(lb, 5):
        ws5.cell(row=i, column=1, value=nom); ws5.cell(row=i, column=2, value=f'=LEN(C{i})-LEN(SUBSTITUTE(C{i},",",""))+1')
        ws5.cell(row=i, column=3, value=", ".join(map(str, ans)))
    ch = BarChart(); ch.type = "bar"; ch.gapWidth = 60
    ch.add_data(Reference(ws5, min_col=2, min_row=4, max_row=4 + len(lb)), titles_from_data=True)
    ch.set_categories(Reference(ws5, min_col=1, min_row=5, max_row=4 + len(lb)))
    styler(ch, "Ballons d'Or : 8 pour Messi, record absolu", legende=False, num="0")
    s = ch.series[0]; remplir(s, GRIS); etiquettes(s, "0")
    s.dPt = [point([k for k, _ in lb].index("Messi"), ORANGE)]
    ch.x_axis.scaling.min = 0
    onglet_graphique(wb, "G4 Ballons", ch)
    for w in wb.worksheets:
        for row in w.iter_rows():
            for cell in row:
                if cell.alignment is None:
                    cell.alignment = Alignment()
    return wb


if __name__ == "__main__":
    for nom, fab, graphes in (
        ("sports-combat", combat, {"G1 Taille": "G1-taille.png", "G2 Croissance": "G2-croissance.png", "G3 Judo": "G3-judo.png"}),
        ("messi", messi, {"G1 Totaux": "G1-totaux.png", "G2 Par match": "G2-par-match.png", "G3 xG": "G3-xg.png", "G4 Ballons": "G4-ballons.png"}),
    ):
        dossier = RACINE / "projets" / nom / "excel"
        dossier.mkdir(parents=True, exist_ok=True)
        f = dossier / f"{nom}.xlsx"
        fab().save(f)
        if RECALC.exists():
            print(f.name, subprocess.run(["python3", str(RECALC), str(f), "120"], capture_output=True, text=True).stdout[:160])
        exporter_graphiques(f, dossier / "graphiques", graphes)
        print(nom, "graphiques exportés")
