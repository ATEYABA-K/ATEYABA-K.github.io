"""
Fichiers pour refaire chaque graphique du portfolio soi-même, sans le site :
  - un CSV par graphique (Excel, Power BI, Google Sheets, Python...)
  - un classeur Excel par étude, avec les mêmes graphiques en natif et de vraies formules
Lancer après les scripts d'analyse :  python outils/reproduire.py
"""
import csv
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference, ScatterChart, Series
from openpyxl.comments import Comment
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

RACINE = Path(__file__).resolve().parent.parent
POLICE = "Arial"
TITRE = Font(name=POLICE, size=14, bold=True)
GRAS = Font(name=POLICE, bold=True)
NORMAL = Font(name=POLICE)
SAISIE = Font(name=POLICE, color="0000FF", bold=True)
JAUNE = PatternFill("solid", fgColor="FFFF00")
ENTETE = PatternFill("solid", fgColor="1F3864")
BLANC = Font(name=POLICE, bold=True, color="FFFFFF")


def ecrire_csv(dossier, nom, entetes, lignes):
    dossier.mkdir(parents=True, exist_ok=True)
    with open(dossier / nom, "w", newline="", encoding="utf-8-sig") as f:  # utf-8-sig : accents OK dans Excel
        w = csv.writer(f, delimiter=";")
        w.writerow(entetes)
        w.writerows(lignes)


def feuille(wb, nom, titre, note, entetes, lignes, formats=None):
    """Crée une feuille : titre, note de source, tableau. Renvoie (ws, première ligne de données)."""
    ws = wb.create_sheet(nom)
    ws["A1"], ws["A1"].font = titre, TITRE
    ws["A2"], ws["A2"].font = note, Font(name=POLICE, italic=True, color="666666")
    for j, e in enumerate(entetes, 1):
        c = ws.cell(row=4, column=j, value=e)
        c.font, c.fill, c.alignment = BLANC, ENTETE, Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(j)].width = max(14, len(e) + 2)
    for i, ligne in enumerate(lignes, 5):
        for j, v in enumerate(ligne, 1):
            c = ws.cell(row=i, column=j, value=v)
            c.font = NORMAL
            if formats and formats.get(j):
                c.number_format = formats[j]
    ws.row_dimensions[4].height = 32
    ws.freeze_panes = "A5"
    return ws, 5


def lisez_moi(wb, titre, lignes):
    ws = wb.active
    ws.title = "Lisez-moi"
    ws["A1"], ws["A1"].font = titre, TITRE
    for i, l in enumerate(lignes, 3):
        ws.cell(row=i, column=1, value=l).font = GRAS if l and not l.startswith(("•", " ")) else NORMAL
    ws.column_dimensions["A"].width = 120


def graphique(ws, chart, ancre, titre, x_titre=None, y_titre=None):
    chart.title, chart.height, chart.width = titre, 9, 18
    if x_titre:
        chart.x_axis.title = x_titre
    if y_titre:
        chart.y_axis.title = y_titre
    chart.x_axis.delete = False
    chart.y_axis.delete = False
    ws.add_chart(chart, ancre)


# =========================================================================
# Étude 1 : MMA × paris sportifs
# =========================================================================
D = RACINE / "projets/ufc-paris"
R = json.loads((D / "resultats.json").read_text())
C = json.loads((D / "combats.json").read_text())
OUT = D / "reproduire"
wb = Workbook()
lisez_moi(wb, "Le bookmaker gagne toujours. Même au MMA. — refaire les graphiques", [
    "Ce classeur contient les données de chaque graphique de l'étude, et le graphique refait dans Excel.",
    "",
    "Les feuilles",
    "• Argent : mises, gains redistribués, part gardée (ARJEL / ANJ, 2010-2025). Colonnes D et E = formules.",
    "• Evenements : nombre d'événements UFC par année (ufcstats.com).",
    "• Calibration : probabilité annoncée par la cote vs taux de victoire réel.",
    "• Rendement : ce que rapporte 1 € misé, par tranche de cote.",
    "• Simulateur : un parieur, 100 paris sur de vrais combats tirés au hasard. Appuyez sur F9 pour relancer. Changez la mise (cellule jaune).",
    "• Combats : les 6 916 combats (cotes décimales, vainqueur) utilisés par le simulateur.",
    "",
    "Refaire un graphique dans Excel",
    "  1. Sélectionnez le tableau (en-têtes compris).  2. Insertion > Graphiques recommandés.  3. Choisissez le type indiqué dans le titre de la feuille.",
    "",
    "Dans Power BI",
    "  Accueil > Obtenir les données > Texte/CSV > choisissez un fichier du dossier reproduire/ (séparateur : point-virgule).",
    "  Puis glissez l'année en Axe et les montants en Valeurs (graphique en colonnes empilées pour la feuille Argent).",
    "",
    "En Python : python graphiques_python.py (dans le même dossier) redessine tous les graphiques en PNG.",
    "",
    "Sources : Ultimate UFC Dataset (github.com/shortlikeafox/ultimate_ufc_dataset), ufcstats.com, rapports ARJEL et ANJ.",
])

# Argent
M = R["marche_francais"]["par_annee"]
lignes = [[int(a), v["mises"], v["garde"]] for a, v in M.items()]
ecrire_csv(OUT, "argent.csv", ["annee", "mises_meur", "garde_meur", "redistribue_meur"],
           [[a, m, g, m - g] for a, m, g in lignes])
ws, r0 = feuille(wb, "Argent", "Paris sportifs en ligne en France (M€) — colonnes empilées",
                 "Sources : ARJEL (2010-2018), ANJ (2020-2025). 2019 : mises déduites (+6 % publié en 2020). 2024 : PBJ révisé par l'ANJ.",
                 ["Année", "Mises (M€)", "Gardé par les opérateurs (M€)", "Revenu aux parieurs (M€)", "Part gardée"],
                 [[str(a), m, g] for a, m, g in lignes], {2: "#,##0", 3: "#,##0"})
n = len(lignes)
for i in range(r0, r0 + n):
    ws[f"D{i}"], ws[f"D{i}"].number_format = f"=B{i}-C{i}", "#,##0"
    ws[f"E{i}"], ws[f"E{i}"].number_format = f"=IFERROR(C{i}/B{i},0)", "0.0%"
fin = r0 + n - 1
ws[f"A{fin + 1}"], ws[f"A{fin + 1}"].font = "Total", GRAS
for col, fmt in (("B", "#,##0"), ("C", "#,##0"), ("D", "#,##0")):
    ws[f"{col}{fin + 1}"], ws[f"{col}{fin + 1}"].number_format = f"=SUM({col}{r0}:{col}{fin})", fmt
ws[f"E{fin + 1}"], ws[f"E{fin + 1}"].number_format = f"=C{fin + 1}/B{fin + 1}", "0.0%"
ws[f"A{fin + 3}"] = "Corrélation mises / argent gardé par les opérateurs :"
ws[f"E{fin + 3}"], ws[f"E{fin + 3}"].number_format = f"=CORREL(B{r0}:B{fin},C{r0}:C{fin})", "0.000"
ch = BarChart(); ch.type, ch.grouping, ch.overlap = "col", "stacked", 100
ch.add_data(Reference(ws, min_col=4, min_row=4, max_row=fin), titles_from_data=True)
ch.add_data(Reference(ws, min_col=3, min_row=4, max_row=fin), titles_from_data=True)
ch.set_categories(Reference(ws, min_col=1, min_row=r0, max_row=fin))
graphique(ws, ch, "G4", "Ce qui est misé : revenu aux parieurs + gardé", "Année", "M€")

# Événements
ev = [[int(a), v] for a, v in R["evenements_ufc_par_annee"].items() if int(a) >= 2001]
ecrire_csv(OUT, "evenements_ufc.csv", ["annee", "evenements"], ev)
ws, r0 = feuille(wb, "Evenements", "Événements UFC par année — histogramme", "Source : ufcstats.com",
                 ["Année", "Événements"], [[str(a), v] for a, v in ev])
ch = BarChart(); ch.add_data(Reference(ws, min_col=2, min_row=4, max_row=r0 + len(ev) - 1), titles_from_data=True)
ch.set_categories(Reference(ws, min_col=1, min_row=r0, max_row=r0 + len(ev) - 1)); ch.legend = None
graphique(ws, ch, "D4", "Événements UFC organisés chaque année")

# Calibration
cal = [[p["annoncee"], p["observee"], p["n"]] for p in R["calibration"] if p["n"] >= 100]
ecrire_csv(OUT, "calibration.csv", ["proba_annoncee", "taux_victoire_reel", "combattants"], cal)
ws, r0 = feuille(wb, "Calibration", "Probabilité annoncée vs taux de victoire réel — nuage de points",
                 "Source : Ultimate UFC Dataset, 2010-2026. Marge du bookmaker retirée.",
                 ["Probabilité annoncée", "Taux de victoire réel", "Combattants", "Cote parfaite (diagonale)"],
                 cal, {1: "0%", 2: "0%"})
fin = r0 + len(cal) - 1
for i in range(r0, fin + 1):
    ws[f"D{i}"], ws[f"D{i}"].number_format = f"=A{i}", "0%"
sc = ScatterChart(); sc.style = 13
xs = Reference(ws, min_col=1, min_row=r0, max_row=fin)
for col, nom in ((2, "Réel"), (4, "Cote parfaite")):
    s = Series(Reference(ws, min_col=col, min_row=r0, max_row=fin), xs, title=nom)
    if col == 2:
        s.marker.symbol, s.graphicalProperties.line.noFill = "circle", True
    sc.series.append(s)
graphique(ws, sc, "F4", "Les bookmakers voient juste", "Probabilité annoncée", "Taux de victoire réel")

# Rendement par cote
rd = [[k, v["rendement"], R["rendement_par_cote_ic95"][k], v["victoires"], v["n"]] for k, v in R["rendement_par_cote"].items()]
ecrire_csv(OUT, "rendement_par_cote.csv", ["tranche_cote", "rendement", "marge_erreur_95", "taux_victoire", "paris"], rd)
ws, r0 = feuille(wb, "Rendement", "Rendement moyen d'un pari, par tranche de cote — histogramme",
                 "Source : Ultimate UFC Dataset. Rendement = (gains / mises) − 1 en misant 1 € sur chaque combattant de la tranche.",
                 ["Tranche de cote", "Rendement moyen", "Marge d'erreur (95 %)", "Taux de victoire", "Paris"],
                 rd, {2: "0.0%", 3: "0.0%", 4: "0%"})
ch = BarChart(); ch.add_data(Reference(ws, min_col=2, min_row=4, max_row=r0 + len(rd) - 1), titles_from_data=True)
ch.set_categories(Reference(ws, min_col=1, min_row=r0, max_row=r0 + len(rd) - 1)); ch.legend = None
graphique(ws, ch, "G4", "Plus la cote est haute, plus on perd", "Cote", "Rendement")

# Combats (pour le simulateur)
ecrire_csv(OUT, "combats.csv", ["cote_rouge", "cote_bleu", "rouge_gagne", "annee"], C)
ws, r0 = feuille(wb, "Combats", "Les combats UFC et leurs cotes (décimales)", "Source : Ultimate UFC Dataset, 2010-2026.",
                 ["Cote rouge", "Cote bleu", "Rouge gagne (1 = oui)", "Année"], C)
NB = len(C)

# Simulateur (formules aléatoires : F9 pour relancer)
ws = wb.create_sheet("Simulateur", 1)
ws["A1"], ws["A1"].font = "Simulateur : un parieur, 100 paris — appuyez sur F9 pour relancer", TITRE
ws["A2"] = "Chaque ligne tire un vrai combat au hasard dans la feuille Combats, choisit un camp au hasard, et applique la vraie cote et le vrai résultat."
ws["A3"], ws["A3"].font = "Mise par pari (€) :", GRAS
ws["C3"], ws["C3"].font, ws["C3"].fill = 10, SAISIE, JAUNE
ws["C3"].comment = Comment("Cellule à modifier : la mise de chaque pari, en euros.", "Alvin")
ws["E3"], ws["E3"].font = "Bilan final :", GRAS
ws["F3"], ws["F3"].number_format, ws["F3"].font = "=H106", '+#,##0 "€";-#,##0 "€"', GRAS
for j, e in enumerate(["Pari n°", "Ligne du combat", "Camp", "Cote", "Gagné ?", "Gain (€)", "", "Cagnotte (€)"], 1):
    if e:
        c = ws.cell(row=5, column=j, value=e); c.font, c.fill = BLANC, ENTETE
ws["A6"], ws["H6"] = 0, 0
for i in range(7, 107):
    ws[f"A{i}"] = i - 6
    ws[f"B{i}"] = f"=RANDBETWEEN(5,{4 + NB})"
    ws[f"C{i}"] = '=IF(RAND()<0.5,"rouge","bleu")'
    ws[f"D{i}"] = f'=IF(C{i}="rouge",INDEX(Combats!A:A,B{i}),INDEX(Combats!B:B,B{i}))'
    ws[f"E{i}"] = f'=IF(C{i}="rouge",INDEX(Combats!C:C,B{i})=1,INDEX(Combats!C:C,B{i})=0)'
    ws[f"F{i}"] = f"=IF(E{i},$C$3*(D{i}-1),-$C$3)"
    ws[f"H{i}"] = f"=H{i - 1}+F{i}"
    ws[f"F{i}"].number_format = ws[f"H{i}"].number_format = "#,##0.00"
for col, w in zip("ABCDEFGH", (9, 16, 9, 9, 10, 11, 3, 14)):
    ws.column_dimensions[col].width = w
lc = LineChart(); lc.add_data(Reference(ws, min_col=8, min_row=5, max_row=106), titles_from_data=True)
lc.set_categories(Reference(ws, min_col=1, min_row=6, max_row=106)); lc.legend = None
graphique(ws, lc, "J5", "La cagnotte du parieur, pari après pari (F9 = nouveau parieur)", "Pari n°", "€")
wb.save(OUT / "ufc-paris.xlsx")

# =========================================================================
# Étude 2 : NBA, le mi-distance
# =========================================================================
D = RACINE / "projets/nba-mi-distance"
R = json.loads((D / "resultats.json").read_text())
T = json.loads((D / "terrain.json").read_text())
OUT = D / "reproduire"
ZONES = ["Sous le panier", "Raquette", "Mi-distance", "3 pts dans le coin", "3 pts dans l'axe"]
S = list(R["part_par_zone"])
wb = Workbook()
lisez_moi(wb, "Le tir que la NBA a arrêté de financer — refaire les graphiques", [
    "Ce classeur contient les données de chaque graphique de l'étude, et le graphique refait dans Excel.",
    "",
    "Les feuilles",
    "• Parts : part des tirs par zone, saison par saison. Les colonnes « grandes zones » sont des formules.",
    "• Rendement : points rapportés par tir, par zone et par saison.",
    "• Decomposition : le calcul « mieux tirer ou tirer au bon endroit », entièrement en formules (SUMPRODUCT).",
    "• Correlation : lien entre part de mi-distance d'une équipe et son rendement, par saison.",
    "• Terrain 2004 / Terrain 2025 : la carte des tirs en carte de chaleur (mise en forme conditionnelle).",
    "",
    "Refaire un graphique dans Excel",
    "  1. Sélectionnez le tableau (en-têtes compris).  2. Insertion > Graphiques recommandés.  3. Choisissez le type indiqué dans le titre de la feuille.",
    "",
    "Dans Power BI",
    "  Obtenir les données > Texte/CSV > terrain.csv : visuel « Nuage de points » avec x en Axe X, y en Axe Y, part en Taille, et saison en Axe de lecture (animation).",
    "",
    "En Python : python graphiques_python.py (dans le même dossier) redessine tous les graphiques en PNG.",
    "",
    "Source : NBA.com, via github.com/DomSamangy/NBA_Shots_04_25 (4,4 millions de tirs, 2003-04 à 2024-25).",
])

# Parts
lignes = [[s] + [R["part_par_zone"][s][z] for z in ZONES] for s in S]
ecrire_csv(OUT, "parts_par_zone.csv", ["saison"] + ZONES, lignes)
ws, r0 = feuille(wb, "Parts", "Part des tirs par zone — courbes", "Source : NBA.com. Lancers francs exclus.",
                 ["Saison"] + ZONES + ["Près du panier", "3 points", "Mi-distance (rappel)"], lignes,
                 {k: "0.0%" for k in range(2, 7)})
fin = r0 + len(S) - 1
for i in range(r0, fin + 1):
    ws[f"G{i}"], ws[f"H{i}"], ws[f"I{i}"] = f"=B{i}+C{i}", f"=E{i}+F{i}", f"=D{i}"
    for col in "GHI":
        ws[f"{col}{i}"].number_format = "0.0%"
lc = LineChart()
for col in (7, 8, 9):
    lc.add_data(Reference(ws, min_col=col, min_row=4, max_row=fin), titles_from_data=True)
lc.set_categories(Reference(ws, min_col=1, min_row=r0, max_row=fin))
graphique(ws, lc, "K4", "Le budget a changé de canal", "Saison", "Part des tirs")

# Rendement
lignes = [[s] + [R["rendement_par_zone"][s][z] for z in ZONES] for s in S]
ecrire_csv(OUT, "rendement_par_zone.csv", ["saison"] + ZONES, lignes)
ws, r0 = feuille(wb, "Rendement", "Points rapportés par tir, par zone — histogramme (première et dernière saison)",
                 "Source : NBA.com. Rendement = points marqués / tirs tentés.", ["Saison"] + ZONES, lignes,
                 {k: "0.00" for k in range(2, 7)})
fin = r0 + len(S) - 1
ch = BarChart()
for row in (r0, fin):
    ch.add_data(Reference(ws, min_col=2, max_col=6, min_row=row), from_rows=True, titles_from_data=False)
ch.series[0].tx = ch.series[1].tx = None
ch.set_categories(Reference(ws, min_col=2, max_col=6, min_row=4))
from openpyxl.chart.series import SeriesLabel
ch.series[0].tx, ch.series[1].tx = SeriesLabel(v=S[0]), SeriesLabel(v=S[-1])
graphique(ws, ch, "H4", "Cinq canaux, cinq rendements", "Zone", "Points par tir")

# Décomposition en formules
ws = wb.create_sheet("Decomposition")
ws["A1"], ws["A1"].font = "Mieux tirer, ou tirer au bon endroit ? (décomposition mix / adresse)", TITRE
ws["A2"] = "Tout est calculé par formules à partir des feuilles Parts et Rendement (première et dernière saison)."
for j, e in enumerate(["Zone", f"Part {S[0]}", f"Part {S[-1]}", f"Rendement {S[0]}", f"Rendement {S[-1]}"], 1):
    c = ws.cell(row=4, column=j, value=e); c.font, c.fill = BLANC, ENTETE
    ws.column_dimensions[get_column_letter(j)].width = 20
for k, z in enumerate(ZONES):
    i, col = 5 + k, get_column_letter(2 + k)
    ws[f"A{i}"] = z
    ws[f"B{i}"], ws[f"C{i}"] = f"=Parts!{col}5", f"=Parts!{col}{4 + len(S)}"
    ws[f"D{i}"], ws[f"E{i}"] = f"=Rendement!{col}5", f"=Rendement!{col}{4 + len(S)}"
    for c2, f2 in (("B", "0.0%"), ("C", "0.0%"), ("D", "0.000"), ("E", "0.000")):
        ws[f"{c2}{i}"].number_format = f2
rows = [("Rendement global, première saison", "=SUMPRODUCT(B5:B9,D5:D9)"),
        ("Rendement global, dernière saison", "=SUMPRODUCT(C5:C9,E5:E9)"),
        ("Hausse totale (point par tir)", "=B13-B12"),
        ("Effet mix : on tire au bon endroit", "=SUMPRODUCT(C5:C9-B5:B9,D5:D9)"),
        ("Effet adresse : on tire mieux", "=SUMPRODUCT(B5:B9,E5:E9-D5:D9)"),
        ("Interaction", "=SUMPRODUCT(C5:C9-B5:B9,E5:E9-D5:D9)"),
        ("Part du gain due au mix", "=B15/B14"),
        ("Contrôle : mix + adresse + interaction − hausse (doit faire 0)", "=B15+B16+B17-B14")]
for k, (lab, f) in enumerate(rows):
    i = 12 + k
    ws[f"A{i}"], ws[f"B{i}"] = lab, f
    ws[f"A{i}"].font = GRAS if k in (2, 3, 6) else NORMAL
    ws[f"B{i}"].number_format = "0.0%" if k == 6 else "0.000"
ws.column_dimensions["A"].width = 62

# Corrélation
lignes = [[s, R["correlation_mi_distance_rendement"][s], R["equipe_avec_le_moins_de_mi_distance"][s]] for s in S]
ecrire_csv(OUT, "correlation_equipes.csv", ["saison", "correlation", "equipe_moins_de_mi_distance"], lignes)
ws, r0 = feuille(wb, "Correlation", "Corrélation part de mi-distance / rendement des équipes — histogramme",
                 "30 équipes par saison. Sous zéro : moins de mi-distance = plus efficace.",
                 ["Saison", "Corrélation", "Équipe avec le moins de mi-distance"], lignes, {2: "0.00"})
ws.column_dimensions["C"].width = 34
ch = BarChart(); ch.add_data(Reference(ws, min_col=2, min_row=4, max_row=r0 + len(S) - 1), titles_from_data=True)
ch.set_categories(Reference(ws, min_col=1, min_row=r0, max_row=r0 + len(S) - 1)); ch.legend = None
graphique(ws, ch, "E4", "Quand tout le monde quitte un canal", "Saison", "Corrélation")

# Terrain : CSV long + deux cartes de chaleur
lignes = []
for s, cells in T["saisons"].items():
    for cy, cx, part, pps in cells:
        lignes.append([f"{int(s) - 1}-{str(s)[2:]}", cx * 2 - 25 + 1, cy * 2 + 1, round(part / 1000, 5), pps])
ecrire_csv(OUT, "terrain.csv", ["saison", "x_pieds", "y_pieds", "part_des_tirs", "points_par_tir"], lignes)
for s in ("2004", "2025"):
    ws = wb.create_sheet(f"Terrain {s}")
    ws["A1"], ws["A1"].font = f"Carte des tirs {int(s) - 1}-{s[2:]} : part des tirs (‰) par case de 2 × 2 pieds", TITRE
    ws["A2"] = "Ligne 1 = au niveau du panier (fond du terrain). Mise en forme conditionnelle : plus c'est rouge, plus on tire de là."
    grid = {(cy, cx): part for cy, cx, part, _ in T["saisons"][s]}
    for cx in range(T["colonnes"]):
        ws.cell(row=4, column=2 + cx, value=cx * 2 - 24).font = GRAS
        ws.column_dimensions[get_column_letter(2 + cx)].width = 5.5
    for cy in range(T["lignes"]):
        ws.cell(row=5 + cy, column=1, value=cy * 2 + 1).font = GRAS
        for cx in range(T["colonnes"]):
            v = grid.get((cy, cx))
            c = ws.cell(row=5 + cy, column=2 + cx, value=round(v, 1) if v else None)
            c.font, c.number_format = Font(name=POLICE, size=8), "0.0"
    rng = f"B5:{get_column_letter(1 + T['colonnes'])}{4 + T['lignes']}"
    ws.conditional_formatting.add(rng, ColorScaleRule(start_type="num", start_value=0, start_color="FFFFFF",
                                                      mid_type="num", mid_value=10, mid_color="F4B183",
                                                      end_type="max", end_color="C0392B"))
wb.save(OUT / "nba-mi-distance.xlsx")
print("ok")
